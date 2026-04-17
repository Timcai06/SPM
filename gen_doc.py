from docx import Document
from docx.shared import Pt

doc = Document()

doc.add_heading("数据库表结构说明文档", 0)

doc.add_heading("1. companies (公司表)", level=1)
table1 = doc.add_table(rows=1, cols=3)
table1.style = "Table Grid"
headers1 = ["字段", "解释", "来源"]
for i, h in enumerate(headers1):
    table1.rows[0].cells[i].text = h

data1 = [
    ["id", "主键", "自增"],
    ["ts_code", "股票代码 (如 000001.SZ)", "外部数据导入"],
    ["company_name", "公司名称", "外部数据导入"],
    ["exchange", "交易所 (SZ/SH)", "外部数据导入"],
    ["industry_l1", "一级行业", "外部数据导入"],
    ["industry_l2", "二级行业", "外部数据导入"],
    ["business_scope", "经营范围", "外部数据导入"],
    ["core_products", "核心产品", "外部数据导入"],
    ["concept_tags", "概念标签 (JSON数组)", "外部数据导入"],
    ["is_active", "是否活跃", "系统维护"],
    ["created_at", "创建时间", "系统自动"],
    ["updated_at", "更新时间", "系统自动"],
]
for row_data in data1:
    row = table1.add_row()
    for i, val in enumerate(row_data):
        row.cells[i].text = val

doc.add_paragraph("数据来源: 外部股票数据API (tushare等)")
doc.add_paragraph()

doc.add_heading("2. structured_events (结构化事件表)", level=1)
table2 = doc.add_table(rows=1, cols=3)
table2.style = "Table Grid"
for i, h in enumerate(headers1):
    table2.rows[0].cells[i].text = h

data2 = [
    ["id", "主键", "自增"],
    ["event_id", "事件唯一标识", "系统生成"],
    ["candidate_id", "关联候选事件ID", "外键 → int_event_candidates"],
    ["event_name", "事件名称", "NLP提取"],
    ["event_date", "事件日期", "NLP提取/解析"],
    ["source", "来源", "原始文档"],
    ["event_subject_type", "主体类型 (公司/人物/行业等)", "分类模型"],
    ["duration_type", "持续类型 (短期/长期)", "分类模型"],
    ["predictability_type", "可预测性 (可预测/不可预测)", "分类模型"],
    ["industry_type", "行业类型", "分类模型"],
    ["sentiment", "情感 (正面/负面/中性)", "分类模型"],
    ["heat_score", "热度评分", "计算得出"],
    ["intensity_score", "强度评分", "计算得出"],
    ["impact_scope", "影响范围", "NLP提取"],
    ["event_summary", "事件摘要", "NLP生成"],
    ["subject_entities", "主体实体 (JSON)", "NLP提取"],
    ["raw_text_ref", "原文引用", "原始文档"],
    ["classification_evidence", "分类证据", "模型输出"],
    ["source_type", "来源类型", "分类"],
    ["source_credibility_score", "来源可信度", "评分模型"],
    ["event_subject_subtype", "主体子类型", "分类模型"],
    ["event_stage", "事件阶段", "分类模型"],
    ["shock_source_type", "冲击来源类型", "分类模型"],
    ["trigger_word_score", "触发词评分", "NLP计算"],
    ["explicitness_score", "明确性评分", "NLP计算"],
    ["uncertainty_score", "不确定性评分", "NLP计算"],
    ["novelty_score", "新颖性评分", "NLP计算"],
    ["event_code", "事件代码", "编码"],
    ["authority_level", "权威级别", "分类模型"],
    ["time_orientation", "时间取向", "分类模型"],
    ["region_scope", "区域范围", "分类模型"],
    ["amount_scale", "金额规模", "NLP提取"],
]
for row_data in data2:
    row = table2.add_row()
    for i, val in enumerate(row_data):
        row.cells[i].text = val

doc.add_paragraph("数据来源: 通过 NLP pipeline 从 raw_documents 提取并分类")
doc.add_paragraph()

doc.add_heading("3. raw_documents (原始文档表)", level=1)
table3 = doc.add_table(rows=1, cols=3)
table3.style = "Table Grid"
for i, h in enumerate(headers1):
    table3.rows[0].cells[i].text = h

data3 = [
    ["id", "主键", "自增"],
    ["source", "来源名称 (如 新浪财经)", "爬虫/导入"],
    ["source_type", "来源类型", "分类"],
    ["title", "标题", "爬虫/导入"],
    ["content", "正文内容", "爬虫/导入"],
    ["publish_time", "发布时间", "爬虫/解析"],
    ["url", "原文URL", "爬虫"],
    ["symbol_or_subject", "关联股票/主体", "爬虫解析"],
    ["content_hash", "内容哈希", "去重计算"],
    ["created_at", "创建时间", "系统自动"],
    ["updated_at", "更新时间", "系统自动"],
]
for row_data in data3:
    row = table3.add_row()
    for i, val in enumerate(row_data):
        row.cells[i].text = val

doc.add_paragraph("数据来源: 爬虫采集 / API导入")
doc.add_paragraph()

doc.add_heading("4. event_company_links (事件公司关联表)", level=1)
table4 = doc.add_table(rows=1, cols=3)
table4.style = "Table Grid"
for i, h in enumerate(headers1):
    table4.rows[0].cells[i].text = h

data4 = [
    ["id", "主键", "自增"],
    ["structured_event_id", "事件ID", "外键 → structured_events"],
    ["company_id", "公司ID", "外键 → companies"],
    ["link_type", "关联类型 (candidate/confirmed等)", "分类/确认"],
    ["relation_path", "关系路径", "NLP提取"],
    ["text_similarity_score", "文本相似度", "计算得出"],
    ["industry_match_score", "行业匹配度", "计算得出"],
    ["chain_position_score", "产业链位置评分", "计算得出"],
    ["event_match_score", "事件匹配度", "计算得出"],
    ["final_link_score", "最终关联评分", "综合计算"],
    ["evidence", "证据 (JSON)", "NLP提取"],
    ["is_manual_override", "是否人工 override", "人工确认"],
    ["created_at", "创建时间", "系统自动"],
    ["updated_at", "更新时间", "系统自动"],
]
for row_data in data4:
    row = table4.add_row()
    for i, val in enumerate(row_data):
        row.cells[i].text = val

doc.add_paragraph("数据来源: 通过匹配算法关联 structured_events 与 companies")

import os

desktop = os.path.expanduser("~/Desktop")
doc.save(os.path.join(desktop, "表结构说明.docx"))
print(f"已保存到: {desktop}/表结构说明.docx")
